# -*- coding: utf-8 -*-
import sys, math
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

import tc_data as D

OUT = sys.argv[1]

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="FF674EA7")
SUB_FILL = PatternFill("solid", fgColor="FFD9D2E9")
INPUT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
TOTAL_FILL = PatternFill("solid", fgColor="FFEFEFEF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

def hdr(ws, row, values, fill=HDR_FILL, color="FFFFFF"):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=FONT, bold=True, color=color, size=10)
        c.fill = fill
        c.alignment = CENTER
        c.border = BORDER

def cell(ws, row, col, value, bold=False, fill=None, align=WRAP_TOP, fmt=None, color=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, bold=bold, size=10, color=color)
    c.alignment = align
    c.border = BORDER
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    return c

def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()

# ============================ 테스트케이스 ============================
ws = wb.active
ws.title = "테스트케이스"
TC_HDR = ["TC ID", "대분류", "중분류", "요구사항 ID", "설계 근거", "실행 구분", "우선순위",
          "사전 조건", "테스트 절차", "테스트 데이터", "기대 결과", "관련 질의",
          "결과", "실측·메모", "결함 ID", "수행일", "수행자", "스크린샷", "비고"]
hdr(ws, 1, TC_HDR)
TC_W = [13, 13, 13, 16, 11, 10, 9, 36, 44, 20, 48, 11, 10, 30, 11, 11, 9, 14, 14]
widths(ws, TC_W)
INPUT_COLS = [13, 14, 15, 16, 17, 18]
for r, tc in enumerate(D.TCS, 2):
    for cidx, v in enumerate(tc, 1):
        a = CENTER if cidx in (1, 2, 3, 4, 5, 6, 7, 12) else WRAP_TOP
        cell(ws, r, cidx, v, align=a)
    for cidx in INPUT_COLS:
        cell(ws, r, cidx, None, fill=INPUT_FILL, align=CENTER if cidx != 14 else WRAP_TOP)
    cell(ws, r, 19, None)
    # 행 높이 추정(줄바꿈·길이 기반)
    lines = 1
    for cidx, w in [(8, TC_W[7]), (9, TC_W[8]), (11, TC_W[10])]:
        txt = str(tc[cidx - 1] or "")
        est = sum(max(1, math.ceil(len(seg) * 1.9 / w)) for seg in txt.split("\n"))
        lines = max(lines, est)
    ws.row_dimensions[r].height = min(15 * lines + 4, 300)
ws.row_dimensions[1].height = 30
LAST = 1 + len(D.TCS)
ws.freeze_panes = "B2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(TC_HDR))}{LAST}"

def dv(ws, formula, rng):
    d = DataValidation(type="list", formula1=formula, allow_blank=True)
    ws.add_data_validation(d); d.add(rng)
dv(ws, '"Pass,Fail,Blocked,N/A"', f"M2:M{LAST}")
dv(ws, '"P1,P2,P3"', f"G2:G{LAST}")
dv(ws, '"요구사항,리스크 기반,추가 검증"', f"E2:E{LAST}")
dv(ws, '"비결제,실결제,장기"', f"F2:F{LAST}")
ws.conditional_formatting.add(f"M2:M{LAST}", CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", fgColor="C6EFCE")))
ws.conditional_formatting.add(f"M2:M{LAST}", CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", fgColor="FFC7CE")))
ws.conditional_formatting.add(f"M2:M{LAST}", CellIsRule(operator="equal", formula=['"Blocked"'], fill=PatternFill("solid", fgColor="FFEB9C")))

T = "'테스트케이스'"
def rng(col): return f"{T}!${col}$2:${col}${LAST}"

# ============================ 요약 ============================
ws = wb.create_sheet("요약", 0)
widths(ws, [16, 11, 8, 8, 8, 9, 9, 9, 9, 60])
cell(ws, 1, 1, f"{D.TITLE} 케이스", bold=True, align=Alignment(vertical="center")).font = Font(name=FONT, bold=True, size=14)
notes = [
    f"- 요구사항 {len(D.REQS)}건(REQ-01~{len(D.REQS)})은 무신사 공개 정책(고객센터 FAQ·이용 안내, {D.DATE} 조회)에서 역도출. 질의사항 {len(D.QUESTIONS)}건(Q01~Q{len(D.QUESTIONS)})은 앱 동작으로 확정하며 기대결과에 'Qn 회신 후 확정' 표기",
    "- 설계 근거: 요구사항(공개 정책에서 직접 도출) / 리스크 기반(정책 외, 금액·주문 상태 영향 큼) / 추가 검증(범위 외 품질 확인)",
    "- 입력: '테스트케이스' 시트 노란 칸. 결과는 Pass/Fail/Blocked/N/A 선택. 집계는 자동 계산",
    "- 실행 구분: 비결제(결제 승인 전까지) / 실결제(소액 1건, 완료 후 즉시 취소) / 장기(48시간 대기 등). 실결제·장기 TC는 묶어서 수행",
]
for i, n in enumerate(notes, 2):
    c = ws.cell(row=i, column=1, value=n); c.font = Font(name=FONT, size=9)
for i in range(1, 6):
    ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=10)

r0 = 7
hdr(ws, r0, ["대분류", "TC 수", "P1", "P2", "P3", "Pass", "Fail", "Blocked", "미완료"])
for i, cat in enumerate(D.CATEGORIES):
    r = r0 + 1 + i
    cell(ws, r, 1, cat, align=CENTER)
    cell(ws, r, 2, f"=COUNTIF({rng('B')},A{r})", align=CENTER)
    for j, p in enumerate(["P1", "P2", "P3"]):
        cell(ws, r, 3 + j, f'=COUNTIFS({rng("B")},$A{r},{rng("G")},"{p}")', align=CENTER)
    for j, s in enumerate(["Pass", "Fail", "Blocked"]):
        cell(ws, r, 6 + j, f'=COUNTIFS({rng("B")},$A{r},{rng("M")},"{s}")', align=CENTER)
    cell(ws, r, 9, f'=B{r}-F{r}-G{r}-H{r}-COUNTIFS({rng("B")},$A{r},{rng("M")},"N/A")', align=CENTER)
rt = r0 + 1 + len(D.CATEGORIES)
cell(ws, rt, 1, "합계", bold=True, fill=TOTAL_FILL, align=CENTER)
for col in range(2, 10):
    L = get_column_letter(col)
    cell(ws, rt, col, f"=SUM({L}{r0+1}:{L}{rt-1})", bold=True, fill=TOTAL_FILL, align=CENTER)
cell(ws, rt + 1, 1, "수행률", bold=True, align=CENTER)
cell(ws, rt + 1, 2, f"=IF(B{rt}=0,0,(F{rt}+G{rt}+H{rt})/B{rt})", align=CENTER, fmt="0.0%")
cell(ws, rt + 1, 3, "Pass율", bold=True, align=CENTER)
cell(ws, rt + 1, 4, f"=IF((F{rt}+G{rt})=0,0,F{rt}/(F{rt}+G{rt}))", align=CENTER, fmt="0.0%")
c = ws.cell(row=rt + 1, column=5, value="수행률 = (Pass+Fail+Blocked)/TC 수, Pass율 = Pass/(Pass+Fail)"); c.font = Font(name=FONT, size=9, italic=True)

r1 = rt + 3
cell(ws, r1, 1, "설계 근거별 분포", bold=True)
hdr(ws, r1 + 1, ["설계 근거", "TC 수", "P1", "P2", "P3", "비율"])
for i, b in enumerate(D.BASIS):
    r = r1 + 2 + i
    cell(ws, r, 1, b, align=CENTER)
    cell(ws, r, 2, f"=COUNTIF({rng('E')},A{r})", align=CENTER)
    for j, p in enumerate(["P1", "P2", "P3"]):
        cell(ws, r, 3 + j, f'=COUNTIFS({rng("E")},$A{r},{rng("G")},"{p}")', align=CENTER)
    cell(ws, r, 6, f"=IF($B${rt}=0,0,B{r}/$B${rt})", align=CENTER, fmt="0.0%")

r2 = r1 + 2 + len(D.BASIS) + 1
cell(ws, r2, 1, "실행 구분별 결과", bold=True)
hdr(ws, r2 + 1, ["실행 구분", "대상 TC", "Pass", "Fail", "Blocked", "미수행"])
for i, e in enumerate(D.EXEC_TYPES):
    r = r2 + 2 + i
    cell(ws, r, 1, e, align=CENTER)
    cell(ws, r, 2, f"=COUNTIF({rng('F')},A{r})", align=CENTER)
    for j, s in enumerate(["Pass", "Fail", "Blocked"]):
        cell(ws, r, 3 + j, f'=COUNTIFS({rng("F")},$A{r},{rng("M")},"{s}")', align=CENTER)
    cell(ws, r, 6, f'=B{r}-C{r}-D{r}-E{r}-COUNTIFS({rng("F")},$A{r},{rng("M")},"N/A")', align=CENTER)

r3 = r2 + 2 + len(D.EXEC_TYPES) + 1
cell(ws, r3, 1, "요구사항 커버리지", bold=True)
hdr(ws, r3 + 1, ["요구사항 ID", "연결 TC 수", "요구사항 기반", "Pass", "Fail", "", "", "", "", "요구사항"])
ws.merge_cells(start_row=r3 + 1, start_column=6, end_row=r3 + 1, end_column=9)
for i, (rid, text, src, imp) in enumerate(D.REQS):
    r = r3 + 2 + i
    cell(ws, r, 1, rid, align=CENTER)
    cell(ws, r, 2, f'=COUNTIF({rng("D")},"*"&A{r}&"*")', align=CENTER)
    cell(ws, r, 3, f'=COUNTIFS({rng("D")},"*"&A{r}&"*",{rng("E")},"요구사항")', align=CENTER)
    cell(ws, r, 4, f'=COUNTIFS({rng("D")},"*"&A{r}&"*",{rng("M")},"Pass")', align=CENTER)
    cell(ws, r, 5, f'=COUNTIFS({rng("D")},"*"&A{r}&"*",{rng("M")},"Fail")', align=CENTER)
    for col in range(6, 10): cell(ws, r, col, None)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
    cell(ws, r, 10, text)
    ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(text) * 1.9 / 60)) + 2

r4 = r3 + 2 + len(D.REQS) + 1
cell(ws, r4, 1, "질의사항 영향", bold=True)
hdr(ws, r4 + 1, ["질의 ID", "영향 TC 수", "", "", "", "", "", "", "", "질의"])
for i, (qid, q, impact, pre) in enumerate(D.QUESTIONS):
    r = r4 + 2 + i
    cell(ws, r, 1, qid, align=CENTER)
    cell(ws, r, 2, f'=COUNTIF({rng("L")},"*"&A{r}&"*")', align=CENTER)
    for col in range(3, 10): cell(ws, r, col, None)
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=9)
    cell(ws, r, 10, q)
    ws.row_dimensions[r].height = 15 * max(1, math.ceil(len(q) * 1.9 / 60)) + 2
c = ws.cell(row=r4 + 2 + len(D.QUESTIONS), column=1, value="※ 질의 TC는 '현재 동작 기록'이 수행 조건이며, 회신(확인) 결과는 '질의사항' 시트에 기록")
c.font = Font(name=FONT, size=9, italic=True)

# ============================ 질의사항 ============================
ws = wb.create_sheet("질의사항")
hdr(ws, 1, ["ID", "질의", "영향", "회신 전 처리", "영향 TC 수", "영향 TC", "확인 결과(앱 동작)", "확인일"])
widths(ws, [7, 55, 22, 40, 10, 34, 45, 11])
for i, (qid, q, impact, pre) in enumerate(D.QUESTIONS, 2):
    linked = ", ".join(t[0].replace("TC-", "") for t in D.TCS if qid in t[11])
    cell(ws, i, 1, qid, align=CENTER); cell(ws, i, 2, q); cell(ws, i, 3, impact); cell(ws, i, 4, pre)
    cell(ws, i, 5, f'=COUNTIF({rng("L")},"*"&A{i}&"*")', align=CENTER)
    cell(ws, i, 6, linked)
    cell(ws, i, 7, None, fill=INPUT_FILL); cell(ws, i, 8, None, fill=INPUT_FILL, align=CENTER)
    ws.row_dimensions[i].height = 15 * max(2, math.ceil(len(q) * 1.9 / 55)) + 2
ws.freeze_panes = "B2"
c = ws.cell(row=len(D.QUESTIONS) + 3, column=1, value=f"※ 영향 TC 목록은 {D.DATE} 설계 시점 기준. TC의 '관련 질의' 열을 바꾸면 '영향 TC 수'만 자동 갱신됨")
c.font = Font(name=FONT, size=9, italic=True)

# ============================ 테스트 데이터 ============================
ws = wb.create_sheet("테스트 데이터")
hdr(ws, 1, ["구분", "ID", "이름", "조건", "사용 TC", "준비 완료", "실제 값(상품명·쿠폰명 등)"])
widths(ws, [8, 8, 26, 48, 30, 10, 40])
for i, row in enumerate(D.TEST_DATA, 2):
    for j, v in enumerate(row, 1):
        cell(ws, i, j, v, align=CENTER if j <= 2 else WRAP_TOP)
    cell(ws, i, 6, None, fill=INPUT_FILL, align=CENTER); cell(ws, i, 7, None, fill=INPUT_FILL)
    ws.row_dimensions[i].height = 32
dv(ws, '"Y,N"', f"F2:F{len(D.TEST_DATA)+1}")
c = ws.cell(row=len(D.TEST_DATA) + 3, column=1, value="※ 금액·조건은 예시. 실제 쿠폰 조건과 브랜드 배송 정책은 수행 시점에 확인해 '실제 값'에 기록")
c.font = Font(name=FONT, size=9, italic=True)

# ============================ 모바일 체크리스트 ============================
ws = wb.create_sheet("모바일 체크리스트")
hdr(ws, 1, ["관점", "확인 항목", "연결 TC", "비고"])
widths(ws, [16, 60, 44, 30])
for i, row in enumerate(D.MOBILE_CHECKLIST, 2):
    for j, v in enumerate(row, 1):
        cell(ws, i, j, v, align=CENTER if j == 1 else WRAP_TOP)
    cell(ws, i, 4, None, fill=INPUT_FILL)
    ws.row_dimensions[i].height = 32
c = ws.cell(row=len(D.MOBILE_CHECKLIST) + 3, column=1, value="※ 앱 테스트 학습용 관점 목록. 각 관점이 최소 1개 TC로 연결되도록 설계")
c.font = Font(name=FONT, size=9, italic=True)

# ============================ 결함 리포트 ============================
ws = wb.create_sheet("결함 리포트")
DEF_HDR = ["결함 ID", "TC ID", "제목", "심각도", "우선순위", "환경(기기/iOS/앱 버전)", "재현 절차", "기대 결과", "실제 결과", "재현율", "첨부", "상태", "등록일", "비고"]
hdr(ws, 1, DEF_HDR)
widths(ws, [10, 13, 30, 8, 8, 20, 40, 30, 30, 8, 14, 9, 11, 16])
example = ["DEF-000", "TC-PNT-004", "(예시) 적립금 1,005원 입력 시 단위 안내 없이 1,005원 그대로 적용됨", "S2", "P1",
           "iPhone 15 / iOS 18.x / 앱 x.y.z", "1. 주문서 적립금 입력 1,005원\n2. 적용 버튼 탭\n3. 최종 결제 금액 확인",
           "10원 단위 절사(1,000원) 또는 단위 안내 후 거부", "1,005원이 그대로 차감되어 최종 금액이 5원 단위로 표시", "3/3",
           "def-000-1.png", "신규", "2026-10-26", "예시 행 — 실제 결함 등록 시 삭제"]
for j, v in enumerate(example, 1):
    cell(ws, 2, j, v, fill=INPUT_FILL, align=CENTER if j in (1, 2, 4, 5, 10, 12, 13) else WRAP_TOP, color="808080")
ws.row_dimensions[2].height = 60
for r in range(3, 23):
    for j in range(1, len(DEF_HDR) + 1):
        cell(ws, r, j, None, fill=INPUT_FILL)
dv(ws, '"S1,S2,S3,S4"', "D2:D22"); dv(ws, '"P1,P2,P3"', "E2:E22"); dv(ws, '"신규,확인,수정중,재검증,종료,보류"', "L2:L22")
legend_r = 25
cell(ws, legend_r, 1, "심각도 기준", bold=True)
for i, (s, d) in enumerate([
    ("S1", "결제 금액 오류, 중복 결제, 주문·데이터 유실, 크래시로 결제 불가"),
    ("S2", "주요 기능(쿠폰·적립금·취소) 동작 불가 또는 정책과 다른 금액 처리"),
    ("S3", "안내 문구·UI 오류, 우회 가능한 기능 이상"),
    ("S4", "개선 제안, 사소한 표기"),
], 1):
    cell(ws, legend_r + i, 1, s, align=CENTER); cell(ws, legend_r + i, 2, d)
    ws.merge_cells(start_row=legend_r + i, start_column=2, end_row=legend_r + i, end_column=7)
ws.freeze_panes = "C2"

# 시트 순서: 요약, 테스트케이스, 테스트 데이터, 질의사항, 모바일 체크리스트, 결함 리포트
order = ["요약", "테스트케이스", "테스트 데이터", "질의사항", "모바일 체크리스트", "결함 리포트"]
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)

# ---- 자체 검증(수식 기대값) ----
cats = Counter(t[1] for t in D.TCS); pri = Counter(t[6] for t in D.TCS); basis = Counter(t[4] for t in D.TCS); ex = Counter(t[5] for t in D.TCS)
print("TC total:", len(D.TCS), "| last row:", LAST)
print("by category:", dict(cats)); print("by priority:", dict(pri)); print("by basis:", dict(basis)); print("by exec:", dict(ex))
ids = [t[0] for t in D.TCS]; assert len(ids) == len(set(ids)), "duplicate TC id"
for t in D.TCS:
    assert t[1] in D.CATEGORIES and t[4] in D.BASIS and t[5] in D.EXEC_TYPES and t[6] in ("P1", "P2", "P3"), t[0]
req_cov = {r[0]: sum(1 for t in D.TCS if r[0] in t[3]) for r in D.REQS}
print("req coverage:", req_cov)
q_cov = {q[0]: sum(1 for t in D.TCS if q[0] in t[11]) for q in D.QUESTIONS}
print("q coverage:", q_cov)
